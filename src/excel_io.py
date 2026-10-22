"""
Excel & Multi-format I/O module for VTU Result Automation.
Provides backward-compatible Excel operations alongside modern multi-format
support (CSV, DOCX, PDF, XLSX, XLS) by delegating to file_io.
"""

import openpyxl
import xlwt
import logging
from logger import setup_logger

# Import modern multi-format support
from file_io import (
    load_usns_from_file,
    export_results,
    export_to_csv,
    export_to_docx,
    export_to_pdf,
    export_to_xlsx,
    export_to_xls
)

logger = setup_logger(__name__)
from config import (
    EXCEL_HEADER_ROW,
    EXCEL_SUBHEADER_ROW,
    EXCEL_USN_COLUMN,
    EXCEL_NAME_COLUMN,
    EXCEL_SUBJECTS_START_COLUMN
)


class MockSheetAdapter:
    """Provides a sheet-like cell interface for non-Excel files (CSV, DOCX, PDF, TXT)"""
    def __init__(self, usn_list):
        self.usn_list = usn_list
        self.max_row = len(usn_list)

    def cell(self, row, column):
        class CellValue:
            def __init__(self, val):
                self.value = val

        idx = row - 1
        if 0 <= idx < len(self.usn_list):
            return CellValue(self.usn_list[idx])
        return CellValue(None)


def load_input_workbook(file_path):
    """
    Load the input file containing USN values.
    Supports .xlsx, .xlsm, .xls, .csv, .docx, .pdf, and .txt files.
    
    Args:
        file_path: Path to the input file
        
    Returns:
        Tuple of (workbook, active_sheet) or (None, None) on error
    """
    import os
    ext = os.path.splitext(file_path)[1].lower()

    if ext in [".xlsx", ".xlsm"]:
        try:
            workbook = openpyxl.load_workbook(file_path)
            sheet = workbook.active
            logger.info(f"Loaded input workbook: {file_path}")
            return workbook, sheet
        except Exception as e:
            logger.error(f"Failed to open Excel file: {e}")
            return None, None

    # For other formats (CSV, DOCX, PDF, TXT), use load_usns_from_file and wrap in adapter
    usns = load_usns_from_file(file_path)
    if not usns:
        return None, None

    adapter = MockSheetAdapter(usns)
    return adapter, adapter


def read_usn(sheet, row_index):
    """
    Read USN value from a specific row in the input sheet
    
    Args:
        sheet: The active worksheet or adapter
        row_index: Row number (1-indexed)
        
    Returns:
        USN string or None if empty
    """
    if hasattr(sheet, "cell"):
        cell = sheet.cell(row=row_index, column=1)
        return cell.value
    return None


def create_output_workbook():
    """
    Create a new output workbook for storing results
    
    Returns:
        Tuple of (workbook, sheet, orange_style)
    """
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Sheet1", cell_overwrite_ok=True)
    workbook.add_sheet("sheet2", cell_overwrite_ok=True)
    orange_style = xlwt.easyxf("pattern: pattern solid, fore_colour orange")
    return workbook, sheet, orange_style


def write_headers(sheet):
    """
    Write static headers to the output Excel sheet
    
    Args:
        sheet: The output worksheet
    """
    if hasattr(sheet, "write"):
        sheet.write(EXCEL_HEADER_ROW, EXCEL_USN_COLUMN, "USN")
        sheet.write(EXCEL_HEADER_ROW, EXCEL_NAME_COLUMN, "NAME")


def write_student_info(sheet, row_index, usn, name):
    """
    Write student USN and name to the output sheet
    
    Args:
        sheet: The output worksheet
        row_index: Row number to write to
        usn: Student USN
        name: Student name
    """
    if hasattr(sheet, "write"):
        sheet.write(row_index + 1, EXCEL_USN_COLUMN, usn)
        sheet.write(row_index + 1, EXCEL_NAME_COLUMN, name)


def write_subject_data(sheet, row_index, subjects, orange_style):
    """
    Write subject data (headers and marks) to the output sheet
    
    Args:
        sheet: The output worksheet
        row_index: Row number to write student data to
        subjects: List of subject dictionaries containing name, ia, see, total, res
        orange_style: Excel style for highlighting result cells
    """
    if not hasattr(sheet, "write"):
        return

    col = EXCEL_SUBJECTS_START_COLUMN
    for sub in subjects:
        sub_name = sub.get("name") or sub.get("code", "")
        sheet.write(EXCEL_HEADER_ROW, col, sub_name)
        sheet.write(EXCEL_SUBHEADER_ROW, col, "IA")
        sheet.write(EXCEL_SUBHEADER_ROW, col + 1, "SEE")
        sheet.write(EXCEL_SUBHEADER_ROW, col + 2, "TOTAL")
        sheet.write(EXCEL_SUBHEADER_ROW, col + 3, "RES")

        try:
            ia_val = int(sub.get("ia", 0))
        except (ValueError, TypeError):
            ia_val = str(sub.get("ia", "-"))

        try:
            see_val = int(sub.get("see", 0))
        except (ValueError, TypeError):
            see_val = str(sub.get("see", "-"))

        try:
            tot_val = int(sub.get("total", 0))
        except (ValueError, TypeError):
            tot_val = str(sub.get("total", "-"))

        sheet.write(row_index + 1, col, ia_val)
        sheet.write(row_index + 1, col + 1, see_val)
        sheet.write(row_index + 1, col + 2, tot_val)
        sheet.write(row_index + 1, col + 3, sub.get("res", "-"), orange_style)
        col += 4


def save_workbook(workbook_or_results, file_path):
    """
    Save the results to the specified file.
    If workbook_or_results is a list of results, uses export_results (supporting CSV, DOCX, PDF, XLSX, XLS).
    If it is an xlwt workbook, saves directly.
    """
    if isinstance(workbook_or_results, list):
        export_results(workbook_or_results, file_path)
    elif hasattr(workbook_or_results, "save"):
        try:
            workbook_or_results.save(file_path)
            logger.info(f"Saved results to {file_path}")
        except Exception as e:
            logger.error(f"Failed to save workbook: {e}")
            raise
    else:
        logger.error(f"Unsupported object passed to save_workbook: {type(workbook_or_results)}")
