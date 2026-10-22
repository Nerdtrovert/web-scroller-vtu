"""
Direct In-Browser Runner for VTU Result Automation with Full Auto-CAPTCHA.
Runs browser in visible or automated mode.
Automatically solves VTU captchas via multi-engine OCR (Apple Vision / Tesseract),
submits, handles retries, and exports department results to Excel.
"""

import os
import sys
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import NoAlertPresentException, NoSuchElementException

from captcha_solver import CaptchaSolver
from file_io import export_results

import platform

PORTAL_URL = "https://results.vtu.ac.in/MJ26cbcs/index.php"
TARGET_USNS = ["1HC24CS058", "1HC24CS088"]
OUTPUT_EXCEL = "vtu_results_output.xlsx"
MAX_CAPTCHA_RETRIES = 5


def get_brave_binary():
    """Detect Brave browser executable across macOS, Windows, and Linux."""
    sys_name = platform.system()
    if sys_name == "Darwin":
        p = "/Applications/Brave Browser.app/Contents/MacOS/Brave Browser"
        if os.path.exists(p):
            return p
    elif sys_name == "Windows":
        win_candidates = [
            os.path.expandvars(r"%PROGRAMFILES%\BraveSoftware\Brave-Browser\Application\brave.exe"),
            os.path.expandvars(r"%PROGRAMFILES(X86)%\BraveSoftware\Brave-Browser\Application\brave.exe"),
            os.path.expandvars(r"%LOCALAPPDATA%\BraveSoftware\Brave-Browser\Application\brave.exe")
        ]
        for p in win_candidates:
            if os.path.exists(p):
                return p
    elif sys_name == "Linux":
        for p in ["/usr/bin/brave-browser", "/usr/bin/brave"]:
            if os.path.exists(p):
                return p
    return None


def setup_driver():
    opts = Options()
    brave_bin = get_brave_binary()
    if brave_bin:
        opts.binary_location = brave_bin
    opts.add_argument("--start-maximized")
    opts.add_experimental_option("detach", True)
    driver = webdriver.Chrome(options=opts)
    return driver


def check_alert(driver):
    try:
        alert = driver.switch_to.alert
        text = alert.text
        alert.accept()
        return text
    except NoAlertPresentException:
        return None
    except Exception:
        return None


def solve_page_captcha(driver):
    """Capture current captcha image from browser and solve it."""
    try:
        captcha_elem = driver.find_element(By.XPATH, '//img[contains(@src, "captcha") or @alt="CAPTCHA code"]')
        img_bytes = captcha_elem.screenshot_as_png
        
        # Save temp captcha for reference
        with open("temp_captcha.png", "wb") as f:
            f.write(img_bytes)
            
        solved = CaptchaSolver.solve(img_bytes)
        return solved
    except Exception as e:
        print(f"[-] Error solving captcha: {e}")
        return None


def get_cell_text(elem):
    """Retrieve text safely using textContent / innerText fallbacks if element is off-screen."""
    txt = elem.text
    if not txt or not txt.strip():
        txt = elem.get_attribute("innerText") or elem.get_attribute("textContent") or ""
    return txt.strip()


def capture_full_result_screenshot(driver, output_path):
    """
    Capture the complete result card without cutting off lower subjects.
    Uses #dataPrint element screenshot or Chrome CDP full-page capture.
    """
    try:
        # 1. Element screenshot of #dataPrint (cleanest, captures entire result card)
        try:
            data_print = driver.find_element(By.ID, "dataPrint")
            data_print.screenshot(output_path)
            print(f"[+] Saved complete result card screenshot: {output_path}")
            return True
        except Exception:
            pass

        # 2. Chrome DevTools Protocol full-page screenshot
        try:
            import base64
            metrics = driver.execute_cdp_cmd("Page.getLayoutMetrics", {})
            content_size = metrics.get("contentSize", {})
            w = int(content_size.get("width", 1280))
            h = int(content_size.get("height", 1600))
            shot = driver.execute_cdp_cmd("Page.captureScreenshot", {
                "captureBeyondViewport": True,
                "clip": {"x": 0, "y": 0, "width": w, "height": h, "scale": 1}
            })
            with open(output_path, "wb") as f:
                f.write(base64.b64decode(shot["data"]))
            print(f"[+] Saved full-page CDP screenshot: {output_path}")
            return True
        except Exception:
            pass

        # 3. Fallback to viewport screenshot
        driver.save_screenshot(output_path)
        return True
    except Exception as e:
        print(f"[-] Screenshot capture error: {e}")
        return False


def scrape_result_data(driver, usn):
    data = {
        "usn": usn,
        "name": "",
        "semester": "",
        "subjects": []
    }

    # Scroll the results table smoothly into view so all subject rows are active & rendered
    try:
        table_elem = driver.find_element(
            By.XPATH, '//*[@id="dataPrint"]//div[contains(@class, "divTable")] | //*[@id="dataPrint"]//table'
        )
        driver.execute_script("arguments[0].scrollIntoView({behavior: 'smooth', block: 'center'});", table_elem)
        time.sleep(0.4)
    except Exception:
        pass

    # 1. Student Name & USN from table
    try:
        rows = driver.find_elements(By.TAG_NAME, "tr")
        for tr in rows:
            tds = tr.find_elements(By.TAG_NAME, "td")
            if len(tds) >= 2:
                lbl = get_cell_text(tds[0])
                val = get_cell_text(tds[-1]).lstrip(":").strip()
                if "Seat Number" in lbl or "USN" in lbl:
                    data["usn"] = val
                elif "Student Name" in lbl:
                    data["name"] = val
    except Exception as e:
        print(f"[{usn}] Error parsing student info from table: {e}")

    # Fallback for name / usn via xpath
    if not data["name"]:
        try:
            name_el = driver.find_element(By.XPATH, '//*[@id="dataPrint"]//table//tr[2]/td[2]')
            data["name"] = get_cell_text(name_el).lstrip(":").strip()
        except Exception:
            pass

    # 2. Extract semester if available
    try:
        body_text = driver.find_element(By.ID, "dataPrint").text
        for line in body_text.splitlines():
            if "Semester" in line:
                data["semester"] = line.strip()
                break
    except Exception:
        pass

    # 3. Extract subjects from divTable
    try:
        row_elems = driver.find_elements(By.XPATH, '//*[@id="dataPrint"]//div[contains(@class, "divTableRow")]')
        if not row_elems:
            row_elems = driver.find_elements(By.XPATH, '//div[contains(@class, "divTableRow")]')

        for r in row_elems:
            cells = [get_cell_text(c) for c in r.find_elements(By.XPATH, './div[contains(@class, "divTableCell")]')]
            if not cells:
                cells = [get_cell_text(c) for c in r.find_elements(By.XPATH, './div') if get_cell_text(c)]

            # Skip header row or empty rows or legend rows
            if not cells or "Subject Code" in cells[0] or "PASS" in cells[0] or "->" in cells[0]:
                continue

            # Standard columns: [Code, Name, IA, SEE, Total, Result, Date]
            if len(cells) >= 6:
                data["subjects"].append({
                    "code": cells[0],
                    "name": cells[1],
                    "ia": cells[2],
                    "see": cells[3],
                    "total": cells[4],
                    "res": cells[5],
                    "date": cells[6] if len(cells) > 6 else ""
                })
    except Exception as e:
        print(f"[{usn}] Error parsing subjects from divTable: {e}")

    # Fallback to standard <table> if divTable was not used
    if not data["subjects"]:
        try:
            for table in driver.find_elements(By.TAG_NAME, "table"):
                for row in table.find_elements(By.TAG_NAME, "tr"):
                    cols = [get_cell_text(col) for col in row.find_elements(By.TAG_NAME, "td")]
                    if len(cols) >= 6 and cols[0] != "Subject Code":
                        data["subjects"].append({
                            "code": cols[0],
                            "name": cols[1],
                            "ia": cols[2],
                            "see": cols[3],
                            "total": cols[4],
                            "res": cols[5],
                            "date": cols[6] if len(cols) > 6 else ""
                        })
        except Exception as e:
            print(f"[{usn}] Table fallback error: {e}")

    return data


def save_to_excel(results_list, output_filename=OUTPUT_EXCEL):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VTU Results"

    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    subheader_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    # Determine unique subjects across all students
    subject_map = {}
    for res in results_list:
        for s in res.get("subjects", []):
            code = s.get("code")
            if code and code not in subject_map:
                subject_map[code] = s.get("name") or code

    subject_list = list(subject_map.keys())

    # Header Row 1
    ws.cell(row=1, column=1, value="USN").font = header_font
    ws.cell(row=1, column=1).fill = header_fill
    ws.cell(row=1, column=2, value="STUDENT NAME").font = header_font
    ws.cell(row=1, column=2).fill = header_fill
    ws.cell(row=1, column=3, value="SEMESTER").font = header_font
    ws.cell(row=1, column=3).fill = header_fill

    col_idx = 4
    for sub in subject_list:
        sub_label = f"{sub} - {subject_map[sub]}"
        ws.cell(row=1, column=col_idx, value=sub_label).font = header_font
        ws.cell(row=1, column=col_idx).fill = header_fill
        ws.merge_cells(start_row=1, start_column=col_idx, end_row=1, end_column=col_idx + 3)
        
        ws.cell(row=2, column=col_idx, value="IA").font = bold_font
        ws.cell(row=2, column=col_idx).fill = subheader_fill
        ws.cell(row=2, column=col_idx + 1, value="SEE").font = bold_font
        ws.cell(row=2, column=col_idx + 1).fill = subheader_fill
        ws.cell(row=2, column=col_idx + 2, value="TOTAL").font = bold_font
        ws.cell(row=2, column=col_idx + 2).fill = subheader_fill
        ws.cell(row=2, column=col_idx + 3, value="RES").font = bold_font
        ws.cell(row=2, column=col_idx + 3).fill = subheader_fill
        col_idx += 4

    # Data Rows
    current_row = 3
    for res in results_list:
        c_usn = ws.cell(row=current_row, column=1, value=res["usn"])
        c_usn.font = bold_font
        c_usn.border = thin_border

        c_name = ws.cell(row=current_row, column=2, value=res["name"])
        c_name.font = regular_font
        c_name.border = thin_border

        c_sem = ws.cell(row=current_row, column=3, value=res.get("semester", ""))
        c_sem.font = regular_font
        c_sem.border = thin_border

        stud_subs = {s.get("code"): s for s in res.get("subjects", [])}

        col_idx = 4
        for sub in subject_list:
            if sub in stud_subs:
                s_data = stud_subs[sub]
                c_ia = ws.cell(row=current_row, column=col_idx, value=s_data.get("ia", "-"))
                c_ia.font = regular_font
                c_ia.border = thin_border

                c_see = ws.cell(row=current_row, column=col_idx + 1, value=s_data.get("see", "-"))
                c_see.font = regular_font
                c_see.border = thin_border

                c_tot = ws.cell(row=current_row, column=col_idx + 2, value=s_data.get("total", "-"))
                c_tot.font = bold_font
                c_tot.border = thin_border

                c_res = ws.cell(row=current_row, column=col_idx + 3, value=s_data.get("res", "-"))
                c_res.font = bold_font
                c_res.border = thin_border
                
                res_val = str(s_data.get("res", "")).upper()
                if res_val == "P":
                    c_res.fill = pass_fill
                elif res_val in ["F", "A"]:
                    c_res.fill = fail_fill
            else:
                for offset in range(4):
                    cell = ws.cell(row=current_row, column=col_idx + offset, value="-")
                    cell.border = thin_border
            col_idx += 4

        current_row += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    wb.save(output_filename)
    print(f"\n[+] Results saved successfully to {os.path.abspath(output_filename)}")


def main():
    print("=" * 65)
    print("VTU Result Automation - Fully Automated Mode (Auto-CAPTCHA)")
    print(f"Portal: {PORTAL_URL}")
    print(f"USNs to process: {TARGET_USNS}")
    print("=" * 65)

    driver = setup_driver()
    results = []

    try:
        for idx, usn in enumerate(TARGET_USNS, 1):
            print(f"\n[{idx}/{len(TARGET_USNS)}] Processing USN: {usn}")
            driver.get(PORTAL_URL)
            time.sleep(1)

            submitted = False
            for attempt in range(1, MAX_CAPTCHA_RETRIES + 1):
                try:
                    usn_input = WebDriverWait(driver, 10).until(
                        EC.presence_of_element_located((By.XPATH, '//input[@name="lns"]'))
                    )
                    usn_input.clear()
                    usn_input.send_keys(usn)

                    captcha_code = solve_page_captcha(driver)
                    print(f"    [Attempt {attempt}/{MAX_CAPTCHA_RETRIES}] Auto-detected CAPTCHA: '{captcha_code}'")

                    if not captcha_code:
                        print("    [-] Failed to solve captcha image. Retrying...")
                        time.sleep(1)
                        continue

                    captcha_input = driver.find_element(By.XPATH, '//input[@name="captchacode"]')
                    captcha_input.clear()
                    captcha_input.send_keys(captcha_code)

                    submit_btn = driver.find_element(By.XPATH, '//*[@id="submit"]')
                    submit_btn.click()

                    time.sleep(1.2)
                    alert = check_alert(driver)
                    if alert:
                        print(f"    [!] Alert from VTU: '{alert}'")
                        if "captcha" in alert.lower():
                            print("    [!] Captcha mismatch detected. Auto-retrying with fresh captcha...")
                            time.sleep(1)
                            continue
                        else:
                            print(f"    [-] Portal notification for {usn}: {alert}. Skipping.")
                            break

                    if "resultpage.php" in driver.current_url:
                        print(f"[+] Result page successfully loaded for {usn}!")
                        submitted = True
                        time.sleep(1)

                        stud_data = scrape_result_data(driver, usn)
                        print(f"[+] Student: {stud_data.get('name', 'N/A')} ({stud_data['usn']})")
                        print(f"[+] Semester: {stud_data.get('semester', 'N/A')}")
                        print(f"[+] Subjects scraped: {len(stud_data['subjects'])}")
                        for s in stud_data['subjects']:
                            print(f"    - {s['code']} ({s['name']}): IA={s['ia']} SEE={s['see']} Total={s['total']} Res={s['res']}")

                        results.append(stud_data)
                        capture_full_result_screenshot(driver, f"result_{usn}.png")
                        break

                except Exception as e:
                    print(f"    [-] Error on attempt {attempt}: {e}")
                    time.sleep(1)

            if not submitted:
                print(f"[-] Could not retrieve result for {usn} after {MAX_CAPTCHA_RETRIES} attempts.")

        if results:
            # Export to all requested formats: Excel (.xlsx), CSV (.csv), Word (.docx), and PDF (.pdf)
            base_name = os.path.splitext(OUTPUT_EXCEL)[0]
            for ext in [".xlsx", ".csv", ".docx", ".pdf"]:
                export_file = f"{base_name}{ext}"
                export_results(results, export_file)
                print(f"[+] Results saved to: {os.path.abspath(export_file)}")

            print("\n" + "=" * 65)
            print("Summary of Scraped Results:")
            for r in results:
                print(f"- {r['usn']} | {r['name']} | {len(r['subjects'])} subjects")
            print("=" * 65)
        else:
            print("[-] No results collected.")

    finally:
        print("\nFinished! Closing browser in 3 seconds...")
        time.sleep(3)
        driver.quit()


if __name__ == "__main__":
    main()
